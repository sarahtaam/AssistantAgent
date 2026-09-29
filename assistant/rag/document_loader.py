"""
assistant/rag/document_loader.py — Multi-format document loading and chunking.

Supports PDF, Markdown, plain text, DOCX, Excel, CSV, and JSON so the
knowledge base can be built from whatever format policy/FAQ documents
actually exist in — a real deployment usually has a mix of these rather
than one tidy format.
"""
from __future__ import annotations
import re
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Dict, Optional
import json

logger = logging.getLogger(__name__)

CHUNK_SIZE = 700
CHUNK_OVERLAP = 150
MIN_CHUNK_LEN = 50

BASE_DIR = Path(__file__).resolve().parent
DOCS_DIR = BASE_DIR / "documents"

SUPPORTED_EXTENSIONS = {
    ".pdf": "pdf",
    ".md": "markdown",
    ".txt": "text",
    ".docx": "docx",
    ".xlsx": "excel",
    ".xls": "excel",
    ".csv": "csv",
    ".json": "json",
}


@dataclass
class Chunk:
    text: str
    source: str
    chunk_id: int
    page: Optional[int] = None
    section: Optional[str] = None
    doc_type: str = "unknown"
    metadata: Dict = field(default_factory=dict)


def _clean(text: str) -> str:
    text = re.sub(r"\r\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _chunk_text(
    text: str, source: str, doc_type: str,
    page: Optional[int] = None, section: Optional[str] = None,
    id_offset: int = 0,
) -> List[Chunk]:
    """Splits text into overlapping chunks, preferring to break on sentence
    boundaries near the target chunk size rather than mid-sentence."""
    text = _clean(text)
    chunks: List[Chunk] = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        raw = text[start:end]
        if end < len(text):
            last = max(raw.rfind(". "), raw.rfind(".\n"), raw.rfind("! "), raw.rfind("? "))
            if last > CHUNK_SIZE // 2:
                end = start + last + 2
                raw = text[start:end]
        if len(raw.strip()) >= MIN_CHUNK_LEN:
            chunks.append(Chunk(
                text=raw.strip(), source=source, chunk_id=id_offset + len(chunks),
                page=page, section=section, doc_type=doc_type,
            ))
        # Once this chunk reaches the end of the text, stop — without this,
        # sections shorter than CHUNK_OVERLAP cause `end - CHUNK_OVERLAP` to
        # fall behind `start`, so the loop re-chunks the same short tail
        # over and over with a 1-character-shrinking window (a 2KB file
        # exploded into 600+ near-duplicate chunks before this fix).
        if end >= len(text):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return chunks


def _load_pdf(filepath: str, source: str) -> List[Chunk]:
    import pypdf  # optional dep (requirements-rag.txt)
    reader = pypdf.PdfReader(filepath)
    chunks: List[Chunk] = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            chunks.extend(_chunk_text(text, source, "pdf", page=i, id_offset=len(chunks)))
    logger.info("PDF %s -> %d chunks (%d pages)", source, len(chunks), len(reader.pages))
    return chunks


def _load_text(filepath: str, source: str) -> List[Chunk]:
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            text = f.read()
    except UnicodeDecodeError:
        with open(filepath, "r", encoding="latin1") as f:
            text = f.read()
    chunks = _chunk_text(text=text, source=source, doc_type="text")
    logger.info("TXT %s -> %d chunks", source, len(chunks))
    return chunks


def _load_markdown(filepath: str, source: str) -> List[Chunk]:
    """Chunks markdown section-by-section (splitting on # / ## / ### headers)
    so each chunk carries its section title as metadata for citation."""
    with open(filepath, "r", encoding="utf-8") as f:
        content = f.read()
    chunks: List[Chunk] = []
    current_section: Optional[str] = None
    current_lines: List[str] = []

    def flush():
        nonlocal current_lines
        text = "\n".join(current_lines).strip()
        if text:
            chunks.extend(_chunk_text(text, source, "markdown", section=current_section, id_offset=len(chunks)))
        current_lines = []

    for line in content.splitlines():
        m = re.match(r"^#{1,3}\s+(.+)$", line.strip())
        if m:
            flush()
            current_section = m.group(1).strip()
        else:
            current_lines.append(line)
    flush()
    logger.info("Markdown %s -> %d chunks", source, len(chunks))
    return chunks


def _load_excel(filepath: str, source: str) -> List[Chunk]:
    import pandas as pd
    chunks = []
    excel_file = pd.ExcelFile(filepath)
    for sheet_name in excel_file.sheet_names:
        df = pd.read_excel(filepath, sheet_name=sheet_name)
        text = df.to_string(index=False)
        if text.strip():
            chunks.extend(_chunk_text(text=text, source=f"{source}:{sheet_name}", doc_type="excel"))
    logger.info("Excel %s -> %d chunks", source, len(chunks))
    return chunks


def _load_csv(filepath: str, source: str) -> List[Chunk]:
    import pandas as pd
    df = pd.read_csv(filepath, encoding="utf-8", sep=None, engine="python")
    text = df.to_string(index=False)
    chunks = _chunk_text(text=text, source=source, doc_type="csv")
    logger.info("CSV %s -> %d chunks", source, len(chunks))
    return chunks


def _load_docx(filepath: str, source: str) -> List[Chunk]:
    from docx import Document  # optional dep (requirements-rag.txt)
    doc = Document(filepath)
    text_parts = []
    text_parts.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text_parts.append(cell.text)
    text = "\n".join(part for part in text_parts if part)
    chunks = _chunk_text(text=text, source=source, doc_type="docx")
    logger.info("DOCX %s -> %d chunks", source, len(chunks))
    return chunks


def _load_json(filepath: str, source: str) -> List[Chunk]:
    with open(filepath, "r", encoding="utf-8") as f:
        data = json.load(f)
    text = json.dumps(data, ensure_ascii=False)
    if len(text) > 20000:
        text = text[:20000]
    return _chunk_text(text=text, source=source, doc_type="json")


class DocumentLoader:
    @classmethod
    def load_all(cls) -> List[Chunk]:
        all_chunks: List[Chunk] = []

        if not DOCS_DIR.exists():
            logger.warning("Documents folder not found: %s", DOCS_DIR)
            return []

        files = sorted(DOCS_DIR.iterdir())
        if not files:
            logger.warning("No documents in: %s", DOCS_DIR)
            return []

        loaders = {
            "pdf": _load_pdf, "markdown": _load_markdown, "text": _load_text,
            "excel": _load_excel, "csv": _load_csv, "json": _load_json, "docx": _load_docx,
        }

        for file_path in files:
            if not file_path.is_file():
                continue
            ext = file_path.suffix.lower()
            if ext not in SUPPORTED_EXTENSIONS:
                logger.warning("Skipped unsupported format: %s", file_path.name)
                continue
            doc_type = SUPPORTED_EXTENSIONS[ext]
            try:
                chunks = loaders[doc_type](str(file_path), file_path.name)
                all_chunks.extend(chunks)
                logger.info("Loaded document: %s", file_path.name)
            except Exception:
                logger.exception("Error loading %s", file_path.name)

        logger.info("Knowledge base loaded: %d chunks", len(all_chunks))
        return all_chunks
